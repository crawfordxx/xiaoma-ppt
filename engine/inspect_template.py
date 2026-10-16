"""拆解 PPT 模板，输出模板画像（profile）。用法：
    python3 inspect_template.py template.pptx --out work/            # 生成 profile.json / profile.md / template-sheet.jpg
    python3 inspect_template.py template.pptx --out work/ --no-render  # 没装 LibreOffice 时只做几何分析

画像包含：设计令牌（主色、辅色、文字色、字体、圆角）、每页角色（cover/toc/section/content/ending/other）、
每页可填的文字槽位（#形状id）、内容页的「内容安全区」、要删掉的示例内容。模型读 profile.md 就能写 deck.json。
"""
import argparse
import collections
import os
import re
import sys

from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from common import (first_size, hls, inch, is_saturated, iter_shapes, luminance, max_size, mix, norm_hex, save_json, shape_text, to_hex, wrap_none, hue_dist)

TOC_RE = re.compile(r"^\s*(目\s*录|contents?|agenda|议\s*程|catalog|目\s*录\s*/?\s*contents?)\s*$", re.I | re.M)
END_RE = re.compile(r"谢\s*谢|感\s*谢|THANKS?|Thanks?|Thank you|THANK YOU|再\s*见|Q\s*&\s*A|敬请|期待.*合作")
NUM_RE = re.compile(r"(PART|Part|CHAPTER|Chapter|SECTION|第)?\s*0?\d{1,2}\s*(章|部分|节)?\s*[.、]?")
DATE_RE = re.compile(r"(19|20)\d{2}\s*[年./\-]|\d{1,2}\s*月|汇报人|演讲人|主讲|日期|DATE", re.I)
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)


# ---------- 主题 ----------
def theme_info(prs):
    master = prs.slide_master
    theme_part = next((r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme")), None)
    fonts = {k: {"latin": "", "ea": ""} for k in ("majorFont", "minorFont")}
    if theme_part is None:
        return {}, fonts
    from lxml import etree
    root = etree.fromstring(theme_part.blob)
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    colors = {}
    scheme = root.find(f".//{a}clrScheme")
    for el in (scheme if scheme is not None else []):
        if not isinstance(el.tag, str) or len(el) == 0:  # 跳过注释
            continue
        name = el.tag.replace(a, "")
        c = el[0]
        try:
            colors[name] = norm_hex(c.get("val") if c.tag.endswith("srgbClr") else c.get("lastClr", "000000"))
        except (ValueError, AttributeError):
            pass
    for kind in ("majorFont", "minorFont"):
        f = root.find(f".//{a}{kind}")
        if f is None:
            continue
        for script in ("latin", "ea"):
            x = f.find(f"{a}{script}")
            fonts[kind][script] = (x.get("typeface") or "") if x is not None else ""
        for s in f.findall(f"{a}font"):
            if s.get("script") == "Hans" and not fonts[kind]["ea"]:
                fonts[kind]["ea"] = s.get("typeface")
    return colors, fonts


SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


def resolve_font(tf, fonts):
    if tf.startswith("+mj"):
        return fonts["majorFont"]["ea" if "ea" in tf else "latin"]
    if tf.startswith("+mn"):
        return fonts["minorFont"]["ea" if "ea" in tf else "latin"]
    return tf


# ---------- 设计令牌 ----------
def _skip(el):
    """p:style（形状默认样式引用主题色）和 p:txStyles（母版默认文字样式）不代表模板的真实用色"""
    p = el.getparent()
    while p is not None:
        if p.tag in (qn("p:style"), qn("p:txStyles"), qn("p:clrMap")):
            return True
        p = p.getparent()
    return False


def collect_style(prs, theme_colors, fonts):
    fill = collections.Counter()
    text = collections.Counter()
    ea, latin, heading = collections.Counter(), collections.Counter(), collections.Counter()
    geoms = collections.Counter()
    parts = [s for s in prs.slides] + list(prs.slide_layouts) + [prs.slide_master]
    for k, part in enumerate(parts):
        el = part._element.find(qn("p:cSld"))
        w = 3 if k < len(prs.slides) else 1  # 页面上的用色权重更高
        for clr in el.iter(qn("a:srgbClr"), qn("a:schemeClr")):
            if _skip(clr):
                continue
            v = clr.get("val")
            if clr.tag == qn("a:schemeClr"):
                v = theme_colors.get(SCHEME_ALIAS.get(v, v))
                if not v:
                    continue
            parent = clr.getparent()
            in_text = parent is not None and parent.getparent() is not None and \
                parent.getparent().tag in (qn("a:rPr"), qn("a:defRPr"))
            if in_text:
                if k < len(prs.slides):
                    text[norm_hex(v)] += 1
                continue
            # 填充色按形状面积加权：大面积的品牌色比小装饰更能代表模板
            sp = clr.getparent()
            while sp is not None and sp.tag not in (qn("p:sp"), qn("p:cxnSp")):
                sp = sp.getparent()
            area = 0.0
            if sp is not None:
                ext = sp.find(".//" + qn("a:ext"))
                if ext is not None and ext.get("cx"):
                    area = int(ext.get("cx")) * int(ext.get("cy")) / (prs.slide_width * prs.slide_height)
            fill[norm_hex(v)] += w * (1 + 20 * min(area, 0.6))
        for r in el.iter(qn("a:rPr"), qn("a:defRPr")):
            if _skip(r):
                continue
            for tag, ctr in ((qn("a:ea"), ea), (qn("a:latin"), latin)):
                f = r.find(tag)
                if f is not None and f.get("typeface"):
                    ctr[resolve_font(f.get("typeface"), fonts)] += 1
            f = r.find(qn("a:ea")) if r.find(qn("a:ea")) is not None else r.find(qn("a:latin"))
            if f is not None and r.get("sz") and int(r.get("sz")) >= 2800:
                heading[resolve_font(f.get("typeface"), fonts)] += 1
        for g in el.iter(qn("a:prstGeom")):
            geoms[g.get("prst")] += 1
    return fill, text, ea, latin, heading, geoms


def brand_dark(c):
    """海军蓝、墨绿这类深色品牌色：亮度低但有明显色相"""
    h, l, s_ = hls(c)
    return 0.07 < l <= 0.18 and s_ > 0.45


def pick_tokens(theme_colors, fonts, style):
    fill, text, ea, latin, heading, geoms = style
    both = fill + text
    sat = [c for c, _ in both.most_common() if is_saturated(c) or brand_dark(c)]
    for k in ("accent1", "accent2"):
        if not sat and theme_colors.get(k):
            sat.append(theme_colors[k])
    primary = sat[0] if sat else "2F5BEA"
    secondary = next((c for c in sat[1:] if hue_dist(c, primary) > 18), None) or mix(primary, "FFFFFF", 0.35)
    from common import contrast
    if contrast(primary, "FFFFFF") < 3 and contrast(secondary, "FFFFFF") >= 4.5 and \
            (is_saturated(secondary) or brand_dark(secondary)):
        primary, secondary = secondary, primary  # 金色这类浅强调色不适合当主色画卡片/图表
    darks = [c for c, _ in text.most_common() if luminance(c) < 0.06]
    dark = darks[0] if darks else (theme_colors.get("dk2") if luminance(theme_colors.get("dk2", "FFFFFF")) < 0.06
                                   else "1F2329")
    font_body = (ea.most_common(1)[0][0] if ea else None) or fonts["minorFont"]["ea"] or \
        (latin.most_common(1)[0][0] if latin else fonts["minorFont"]["latin"])
    font_head = heading.most_common(1)[0][0] if heading else (fonts["majorFont"]["ea"] or font_body)
    rounded = geoms.get("roundRect", 0) + geoms.get("ellipse", 0)
    square = geoms.get("rect", 0)
    return {
        "primary": primary, "secondary": secondary, "dark": dark,
        "body": mix(dark, "FFFFFF", 0.18), "muted": mix(dark, "FFFFFF", 0.45),
        "font": font_body, "font_heading": font_head,
        "font_latin": latin.most_common(1)[0][0] if latin else fonts["minorFont"]["latin"],
        "radius": 0.12 if rounded > square * 0.5 else 0.0,
        "palette": [c for c, _ in both.most_common(12)],
    }


# ---------- 每页分析 ----------
def ph_size(shape, slide):
    """占位符没有显式字号时，去版式和母版里找"""
    s = first_size(shape, 0)
    if s:
        return s
    try:
        idx = shape.placeholder_format.idx
        for lp in slide.slide_layout.placeholders:
            if lp.placeholder_format.idx == idx:
                s = first_size(lp, 0)
                if s:
                    return s
        typ = shape.placeholder_format.type
        tx = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
        st = tx.find(qn("p:titleStyle") if typ in TITLE_TYPES else qn("p:bodyStyle"))
        d = st.find(".//" + qn("a:defRPr"))
        if d is not None and d.get("sz"):
            return int(d.get("sz")) / 100
    except Exception:
        pass
    return 18.0


def is_number(t):
    return bool(t) and len(t) <= 12 and bool(NUM_RE.fullmatch(t.strip()))


def info(shape, box, slide, nested):
    t = shape_text(shape)
    ph = shape.is_placeholder
    size = ph_size(shape, slide) if ph else (max_size(shape) if t else first_size(shape))
    algn = None
    p = shape._element.find(".//" + qn("a:pPr"))
    if p is not None:
        algn = p.get("algn")
    return {
        "key": f"#{shape.shape_id}", "name": shape.name, "text": t, "box": [round(v, 2) for v in box],
        "size": round(size, 1), "wrap": not wrap_none(shape), "align": algn or "l",
        "placeholder": str(shape.placeholder_format.type).split(".")[-1].split(" ")[0] if ph else None,
        "nested": nested, "shape": shape,
    }


def classify(i, n, texts, W, H):
    all_text = "\n".join(t["text"] for t in texts)
    nums = [t for t in texts if is_number(t["text"])]
    if TOC_RE.search(all_text) or (i <= 2 and len(nums) >= 3):
        return "toc"
    if i >= n - 2 and i > 0 and END_RE.search(all_text):
        return "ending"
    if i == n - 1 and n >= 4 and sum(len(t["text"]) for t in texts) < 40:  # 只有 logo/口号的结尾页
        return "ending"
    if i == 0:
        return "cover"
    titles = [t for t in texts if not is_number(t["text"])]
    if any(t["placeholder"] in ("TITLE", "CENTER_TITLE") for t in texts) and \
            any(t["placeholder"] == "TITLE" and t["box"][1] < H * 0.3 for t in texts):
        return "content"
    big = max(titles, key=lambda t: t["size"], default=None)
    has_num = any(is_number(t["text"]) for t in texts)
    chars = sum(len(t["text"]) for t in texts)
    if big and big["box"][1] < H * 0.22 and big["size"] < 34:
        return "content"
    if (has_num or (big and big["size"] >= 28)) and chars < 80:
        return "section"
    return "content"


def reading_order(items):
    return sorted(items, key=lambda t: (round((t["box"][1] + t["box"][3] / 2) / 0.45), t["box"][0]))


def assign_slots(role, texts, W, H):
    slots = {}
    nums = [t for t in texts if is_number(t["text"])]
    words = [t for t in texts if not is_number(t["text"])]
    by_size = sorted(words, key=lambda t: (-t["size"], -t["box"][2] * t["box"][3]))
    title_ph = [t for t in words if t["placeholder"] in ("TITLE", "CENTER_TITLE")]
    if role == "toc":
        heads = [t for t in words if TOC_RE.search(t["text"])]
        rest = [t for t in words if t not in heads]
        # 目录条目字号相同：取出现最多的字号为条目，其余（如「目录」旁的英文副标）归入标题
        common = collections.Counter(t["size"] for t in rest).most_common(1)
        items = reading_order([t for t in rest if common and t["size"] == common[0][0]])
        others = [t for t in rest if t not in items]
        subs = []
        for it in items:  # 条目正下方、字号更小的文字 = 条目副标题
            b = it["box"]
            cand = [t for t in others if t["size"] < it["size"] and 0 <= t["box"][1] - (b[1] + b[3]) < 0.6
                    and abs(t["box"][0] - b[0]) < 0.3]
            subs.append(cand[0] if cand else None)
        if any(subs):
            slots["item_subs"] = [t["key"] if t else None for t in subs]
            others = [t for t in others if t not in subs]
        heads += others
        slots["heading"] = [t["key"] for t in heads]
        slots["items"] = [t["key"] for t in items]
        slots["numbers"] = [t["key"] for t in reading_order(nums)]
    elif role in ("cover", "ending", "section"):
        main = title_ph[0] if title_ph else (by_size[0] if by_size else None)
        if main:
            slots["title"] = main["key"]
        rest = [t for t in by_size if t is not main]
        dates = [t for t in rest if DATE_RE.search(t["text"])]
        slots["subtitle"] = [t["key"] for t in rest if t not in dates]
        if dates:
            slots["meta"] = [t["key"] for t in dates]
        if nums:
            slots["number"] = nums[0]["key"]
    else:
        top = [t for t in words if t["box"][1] < H * 0.3]
        main = (title_ph[0] if title_ph else None) or (max(top, key=lambda t: (t["size"], t["box"][2]))
                                                         if top else None)
        if main:
            slots["title"] = main["key"]
        hn = [t for t in nums if t["box"][1] < H * 0.3]
        if hn:
            slots["number"] = hn[0]["key"]
        if main:
            mb = main["box"][1] + main["box"][3]
            slots["kicker"] = [t["key"] for t in top if t is not main and t["box"][1] < mb + 0.3]
    return slots


def header_footer(slide_shapes, slots, W, H, layout_boxes):
    title = next((s for s in slide_shapes if s["key"] == slots.get("title")), None)
    head_keys = {slots.get("title"), slots.get("number")} | set(slots.get("kicker", []))
    if title:
        tb = title["box"]
        cut = tb[1] + tb[3]
        # 与标题同一水平带内的装饰（色块、线条、编号）都算页眉
        for s in slide_shapes:
            b = s["box"]
            if s["key"] in head_keys or (b[1] < cut + 0.15 and b[3] < H * 0.3 and b[1] + b[3] < H * 0.32):
                cut = max(cut, b[1] + b[3])
    else:
        cut = H * 0.18
    foot = H - 0.4
    for b in [s["box"] for s in slide_shapes] + layout_boxes:
        if b[1] > H * 0.84 and b[3] < H * 0.16 and b[2] < W * 0.98:
            foot = min(foot, b[1] - 0.12)
    return cut + 0.25, foot


def layout_boxes(slide, W, H):
    """版式和母版上的非占位符形状（装饰、logo、页脚），整页大的背景图除外"""
    boxes = []
    layout = slide.slide_layout
    srcs = [layout.shapes]
    if layout._element.get("showMasterSp") != "0":
        srcs.append(layout.slide_master.shapes)
    for shapes in srcs:
        for s, b, _ in iter_shapes(shapes):
            if s.is_placeholder or b[2] * b[3] > W * H * 0.6 or s.shape_type == 6:
                continue
            boxes.append(b)
    return boxes


def free_rect(occupied, region, step=0.1):
    """region 内不与 occupied 相交的最大矩形（直方图法），返回 (x,y,w,h)"""
    x0, y0, x1, y1 = region
    cols, rows = max(1, int((x1 - x0) / step)), max(1, int((y1 - y0) / step))
    grid = [[True] * cols for _ in range(rows)]
    for (bx, by, bw, bh) in occupied:
        c0, c1 = int((bx - x0) / step), int((bx + bw - x0) / step + 0.999)
        r0, r1 = int((by - y0) / step), int((by + bh - y0) / step + 0.999)
        for r in range(max(0, r0), min(rows, r1)):
            for c in range(max(0, c0), min(cols, c1)):
                grid[r][c] = False
    return _max_rect(grid, x0, y0, step)


def _max_rect(grid, x0, y0, step):
    rows, cols = len(grid), len(grid[0])
    hist = [0] * cols
    best = (0, 0, 0, 0, 0)
    for r in range(rows):
        for c in range(cols):
            hist[c] = hist[c] + 1 if grid[r][c] else 0
        stack = []
        for c in range(cols + 1):
            h = hist[c] if c < cols else 0
            start = c
            while stack and stack[-1][1] >= h:
                s, sh = stack.pop()
                area = sh * (c - s)
                if area > best[0]:
                    best = (area, s, r - sh + 1, c - s, sh)
                start = s
            stack.append((start, h))
    _, c, r, w, h = best
    return (round(x0 + c * step, 2), round(y0 + r * step, 2), round(w * step, 2), round(h * step, 2))


def busy_boxes(img_path, W, H, region, step=0.1):
    """从渲染图里找出背景上的「画」（吉祥物、插画、山水），作为占用区域。
    按 0.1 英寸格子统计边缘密度：细线、网格、纹理密度低不算占用，大块图形和文字密度高才算。"""
    im = Image.open(img_path).convert("L")
    edges = im.filter(ImageFilter.FIND_EDGES).point(lambda v: 255 if v > 24 else 0)
    cols, rows = int(W / step), int(H / step)
    blk = 3  # 边缘密度按 0.3 英寸的大格子算：细线、网格占比低，文字和插画占比高
    coarse = edges.resize((max(1, cols // blk), max(1, rows // blk)), Image.BOX)
    small = coarse.resize((cols, rows), Image.NEAREST)
    # 大块平涂图形（山体、色块插画）边缘少：和「同一行、同一列的中位色」都差得多才算（渐变背景不会误报）
    col = Image.open(img_path).convert("RGB").resize((cols, rows), Image.BOX)
    cp = col.load()

    def med(vals):
        return tuple(sorted(v[k] for v in vals)[len(vals) // 2] for k in range(3))

    rmed = [med([cp[c, r] for c in range(cols)]) for r in range(rows)]
    cmed = [med([cp[c, r] for r in range(rows)]) for c in range(cols)]

    def dist(p, q):
        return sum(abs(p[k] - q[k]) for k in range(3))

    px = small.load()
    busy = set()
    for r in range(rows):
        for c in range(cols):
            v = cp[c, r]
            if px[c, r] > 64 or min(dist(v, rmed[r]), dist(v, cmed[c])) > 30:
                busy.add((c, r))
    grown = {(c + dc, r + dr) for c, r in busy for dc in (-1, 0, 1) for dr in (-1, 0, 1)}
    x0, y0, x1, y1 = region
    return [(c * step, r * step, step, step) for c, r in grown
            if 0 <= c < cols and 0 <= r < rows and x0 <= c * step < x1 and y0 <= r * step < y1]


def sample_bg(img_path, W, H, rect):
    im = Image.open(img_path).convert("RGB")
    sx, sy = im.width / W, im.height / H
    x, y, w, h = rect
    x0, y0 = int(x * sx), int(y * sy)
    crop = im.crop((x0, y0, max(x0 + 1, int((x + w) * sx)), max(y0 + 1, int((y + h) * sy)))).resize((40, 24))
    c = collections.Counter(crop.getdata()).most_common(1)[0][0]
    return to_hex(c)


def strip_samples(prs, plan):
    """删掉内容页的示例内容，供渲染「干净」的版面"""
    for i, slide in enumerate(prs.slides):
        for key in plan.get(i, []):
            for s in list(slide.shapes):
                if f"#{s.shape_id}" == key:
                    s._element.getparent().remove(s._element)


def analyze(path, out, do_render=True):
    prs = Presentation(path)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    theme_colors, fonts = theme_info(prs)
    tokens = pick_tokens(theme_colors, fonts, collect_style(prs, theme_colors, fonts))
    slides = []
    n = len(prs.slides)
    # 在 ≥3 页上重复出现的文字是品牌名/页脚（如「睿信资本」），保留原样，不当作可填槽位
    seen = collections.Counter()
    for slide in prs.slides:
        seen.update({shape_text(s) for s, _, _ in iter_shapes(slide.shapes) if s.has_text_frame and shape_text(s)})
    from build import SAMPLE_RE
    brand = {t for t, c in seen.items() if c >= 3 and len(t) <= 30 and not NUM_RE.fullmatch(t.strip())
             and not SAMPLE_RE.search(t)}
    for i, slide in enumerate(prs.slides):
        shapes, texts = [], []
        for s, b, nested in iter_shapes(slide.shapes):
            d = info(s, b, slide, nested)
            shapes.append(d)
            if s.has_text_frame and (d["text"] or s.is_placeholder):
                texts.append(d)
        role = classify(i, n, texts, W, H)
        slots = assign_slots(role, [t for t in texts if t["text"] not in brand], W, H)
        lb = layout_boxes(slide, W, H)
        entry = {"index": i, "role": role, "layout": slide.slide_layout.name, "slots": slots,
                 "texts": [{k: v for k, v in t.items() if k != "shape"} for t in texts]}
        if role == "content":
            top, foot = header_footer(shapes, slots, W, H, lb)
            left = min([shapes_by(shapes, slots.get("title"))["box"][0]] if slots.get("title") else [0.6])
            left = max(0.45, min(left, 1.2))
            region = (left, top, W - left, foot)
            samples = []
            for d in shapes:
                if d["nested"]:
                    continue
                b = d["box"]
                if b[2] * b[3] > W * H * 0.6:
                    continue
                inside = b[1] >= top - 0.3 and b[1] + b[3] <= foot + 0.3
                edge_pic = d["shape"].shape_type == 13 and (b[0] < 0.1 or b[0] + b[2] > W - 0.1)
                if inside and not edge_pic and d["key"] not in (slots.get("title"), slots.get("number")):
                    samples.append(d["key"])
            occ = lb + [d["box"] for d in shapes if not d["nested"] and d["key"] not in samples
                        and d["box"][2] * d["box"][3] < W * H * 0.6]
            entry.update(region=[round(v, 2) for v in region], samples=samples, _occ=occ)
        slides.append(entry)

    files = []
    if do_render:
        from render import render, sheets
        stripped = os.path.join(out, "_stripped.pptx")
        p2 = Presentation(path)
        strip_samples(p2, {s["index"]: s.get("samples", []) for s in slides})
        for sl in p2.slides:  # 隐藏页 LibreOffice 不导出，会让图和页错位
            sl._element.attrib.pop("show", None)
        p2.save(stripped)
        files = render(stripped, os.path.join(out, "template-render"), dpi=60)
        os.remove(stripped)
        if len(files) != n:
            print(f"⚠️ 渲染出 {len(files)} 页、模板有 {n} 页，跳过像素分析", file=sys.stderr)
            files = []
    for s in slides:
        img = files[s["index"]] if files else None
        if s["role"] == "content":
            reg = s["region"]
            occ = s.pop("_occ")
            if img:
                occ = occ + busy_boxes(img, W, H, (reg[0], reg[1], reg[2], reg[3]))
            a = free_rect(occ, reg)
            if a[2] < 0.3 or a[3] < 0.3:  # 整页都是图：退回到几何区域，交给模型判断
                a = (reg[0], reg[1], reg[2] - reg[0], reg[3] - reg[1])
                s["warning"] = "背景图占满了内容区，安全区按几何位置给出，请看渲染图确认"
            pad = 0.05
            s["area"] = [round(a[0] + pad, 2), round(a[1], 2), round(a[2] - 2 * pad, 2), round(a[3], 2)]
            if (a[2] < W * 0.45 or a[3] < H * 0.35) and "warning" not in s:
                s["warning"] = "内容安全区偏小，这一页不适合放大段内容"
        rect = s.get("area") or [W * 0.3, H * 0.3, W * 0.4, H * 0.4]
        if img:
            s["bg"] = sample_bg(img, W, H, rect)
        else:
            s["bg"] = "FFFFFF"
        s["dark"] = luminance(s["bg"]) < 0.22
    prof = {"template": os.path.relpath(os.path.abspath(path), os.path.abspath(out)), "size": [round(W, 3), round(H, 3)], "tokens": tokens,
            "theme_colors": theme_colors, "theme_fonts": fonts, "slides": slides}
    content_bgs = [s["bg"] for s in slides if s["role"] == "content" and not s["dark"]]
    prof["tokens"]["surface"] = "FFFFFF" if not content_bgs or luminance(content_bgs[0]) < 0.9 else \
        mix(content_bgs[0], tokens["primary"], 0.04)
    if files:
        labels = [f'{s["role"]}' for s in slides]
        prof["sheets"] = [os.path.relpath(p, out) for p in
                          sheets(files, os.path.join(out, "template-render"), labels=labels)]
    return prof


def shapes_by(shapes, key):
    return next(s for s in shapes if s["key"] == key)


def capacity(t, W):
    """估算一个槽位一行能放多少个汉字"""
    w = t["box"][2] if t["wrap"] else W * 0.8
    return max(1, int(w * 72 / max(6, t["size"])))


def write_md(prof, path):
    W, H = prof["size"]
    tk = prof["tokens"]
    L = [f"# 模板画像：{os.path.basename(prof['template'])}", "",
         f"- 页面：{W:.2f} × {H:.2f} 英寸，共 {len(prof['slides'])} 页",
         f"- 主色 primary `#{tk['primary']}`，辅色 secondary `#{tk['secondary']}`，标题文字 dark `#{tk['dark']}`，"
         f"正文 body `#{tk['body']}`，弱化 muted `#{tk['muted']}`，卡片底 surface `#{tk['surface']}`",
         f"- 字体：正文 `{tk['font']}`，标题 `{tk['font_heading']}`，西文 `{tk['font_latin']}`；"
         f"圆角 radius={tk['radius']}（0 = 直角风格）",
         f"- 模板高频色：{' '.join('#' + c for c in tk['palette'])}", ""]
    if prof.get("sheets"):
        L += ["模板缩略图（每格右下角是页码 · 角色，路径相对本文件）：" + "，".join(prof["sheets"]), ""]
    by_role = collections.defaultdict(list)
    for s in prof["slides"]:
        by_role[s["role"]].append(s["index"])
    L += ["## 角色总览（deck.json 的 from 填这里的页码，从 0 开始）", ""]
    for r in ("cover", "toc", "section", "content", "ending", "other"):
        if by_role.get(r):
            L.append(f"- **{r}**：{by_role[r]}")
    L.append("")
    for s in prof["slides"]:
        L.append(f"## 第 {s['index']} 页 · {s['role']}（版式：{s['layout']}，底色 #{s['bg']}"
                 f"{'，深色底' if s['dark'] else ''}）")
        for k, v in s["slots"].items():
            L.append(f"- 槽位 `{k}` → {v}")
        if s.get("area"):
            a = s["area"]
            L.append(f"- 内容安全区 area = [{a[0]}, {a[1]}, {a[2]}, {a[3]}]（x, y, 宽, 高，英寸）")
        if s.get("samples"):
            L.append(f"- 构建时自动删除的示例内容：{', '.join(s['samples'])}")
        if s.get("warning"):
            L.append(f"- ⚠️ {s['warning']}")
        L.append("")
        L.append("| 形状 | 文字 | 字号 | 换行 | 一行约容纳 |")
        L.append("|---|---|---|---|---|")
        for t in s["texts"]:
            txt = (t["text"] or "（空占位符）").replace("\n", " / ").replace("|", "｜")
            txt = txt[:40] + ("…" if len(txt) > 40 else "")
            L.append(f"| `{t['key']}` {t['name']} | {txt} | {t['size']} | {'是' if t['wrap'] else '否（构建时自动加宽）'} "
                     f"| {capacity(t, W)} 字 |")
        L.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("--out", default="work")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    prof = analyze(a.template, a.out, not a.no_render)
    save_json(os.path.join(a.out, "profile.json"), prof)
    write_md(prof, os.path.join(a.out, "profile.md"))
    roles = collections.Counter(s["role"] for s in prof["slides"])
    print(f"模板 {len(prof['slides'])} 页：" + "，".join(f"{k}×{v}" for k, v in roles.items()))
    print(f"主色 #{prof['tokens']['primary']}  字体 {prof['tokens']['font']}")
    print("→", os.path.join(a.out, "profile.md"))
    if not any(s["role"] == "content" for s in prof["slides"]):
        print("⚠️ 没识别出内容页：请在 deck.json 里用 area 手动指定内容区域", file=sys.stderr)


if __name__ == "__main__":
    main()
